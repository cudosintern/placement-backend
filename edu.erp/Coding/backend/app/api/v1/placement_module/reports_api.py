from fastapi import APIRouter, Depends, Query, Header, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import text, func, case
from typing import Optional, List
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.core.database import get_db
from app.db.models import IEMStudents, IEMSDepartment, IEMSAcademicBatch
from app.db.placement_models import (
    PLMStudentProfile, PlacementOffer, PlacementDrive, PlacementCompany, PlacementApplication,
    PlacementShortlist, PlacementWaitlist
)
from app.utils.auth_helper import get_current_user
from app.utils.http_return_helper import returnSuccess, returnException

router = APIRouter()

# -------------------------------------------------------------
# 1. Helper to fetch reports query
# -------------------------------------------------------------
def get_student_placement_query(
    db: Session,
    dept_id: Optional[int] = None,
    batch_id: Optional[int] = None,
    placement_status: Optional[str] = None,
    company_id: Optional[int] = None
):
    status_expr = case(
        (PlacementOffer.offer_id.isnot(None), "PLACED"),
        else_="UNPLACED"
    ).label("status")

    query = db.query(
        IEMStudents.usno,
        IEMStudents.name,
        IEMSDepartment.dept_acronym.label("branch"),
        PLMStudentProfile.current_cgpa.label("cgpa"),
        status_expr,
        PlacementCompany.company_name.label("company"),
        PlacementOffer.role.label("designation"),
        PlacementOffer.ctc.label("ctc")
    ).join(
        PLMStudentProfile, IEMStudents.student_id == PLMStudentProfile.student_id
    ).join(
        IEMSDepartment, IEMStudents.department_id == IEMSDepartment.dept_id
    ).join(
        IEMSAcademicBatch, IEMStudents.academic_batch_id == IEMSAcademicBatch.academic_batch_id
    ).outerjoin(
        PlacementOffer, (PLMStudentProfile.profile_id == PlacementOffer.profile_id) & (PlacementOffer.status == "ACCEPTED")
    ).outerjoin(
        PlacementDrive, PlacementOffer.drive_id == PlacementDrive.drive_id
    ).outerjoin(
        PlacementCompany, PlacementDrive.company_id == PlacementCompany.company_id
    )

    if dept_id:
        query = query.filter(IEMStudents.department_id == dept_id)
    if batch_id:
        query = query.filter(IEMStudents.academic_batch_id == batch_id)
    if placement_status and placement_status != "All":
        query = query.filter(status_expr == placement_status)
    if company_id:
        query = query.filter(PlacementDrive.company_id == company_id)

    return query

# -------------------------------------------------------------
# 2. Get Student Placement Status Report Endpoint
# -------------------------------------------------------------
@router.get("/student-placement")
def get_student_placement_report(
    dept_id: Optional[int] = Query(None),
    batch_id: Optional[int] = Query(None),
    placement_status: Optional[str] = Query(None),
    company_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    try:
        query = get_student_placement_query(db, dept_id, batch_id, placement_status, company_id)
        results = query.all()
        
        report_data = []
        for r in results:
            report_data.append({
                "usn": r.usno,
                "name": r.name,
                "branch": r.branch,
                "cgpa": float(r.cgpa) if r.cgpa else 0.0,
                "company": r.company or "-",
                "designation": r.designation or "-",
                "ctc": f"{float(r.ctc)} LPA" if r.ctc else "-",
                "status": r.status
            })
            
        return returnSuccess(report_data, "Student placement report fetched successfully")
    except Exception as e:
        return returnException(str(e))

# -------------------------------------------------------------
# 3. Export Excel Endpoint
# -------------------------------------------------------------
@router.get("/student-placement/export-excel")
def export_student_placement_excel(
    dept_id: Optional[int] = Query(None),
    batch_id: Optional[int] = Query(None),
    placement_status: Optional[str] = Query(None),
    company_id: Optional[int] = Query(None),
    db: Session = Depends(get_db)
):
    try:
        query = get_student_placement_query(db, dept_id, batch_id, placement_status, company_id)
        results = query.all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Placement Report"

        # Headers
        headers = ["USN", "Name", "Branch", "CGPA", "Company", "Designation", "CTC", "Status"]
        ws.append(headers)

        # Style headers
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Rows
        for r in results:
            ws.append([
                r.usno,
                r.name,
                r.branch,
                float(r.cgpa) if r.cgpa else 0.0,
                r.company or "-",
                r.designation or "-",
                f"{float(r.ctc)} LPA" if r.ctc else "-",
                r.status
            ])

        # Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)

        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=Student_Placement_Report.xlsx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# -------------------------------------------------------------
# 4. Export PDF Endpoint
# -------------------------------------------------------------
@router.get("/student-placement/export-pdf")
def export_student_placement_pdf(
    dept_id: Optional[int] = Query(None),
    batch_id: Optional[int] = Query(None),
    placement_status: Optional[str] = Query(None),
    company_id: Optional[int] = Query(None),
    db: Session = Depends(get_db)
):
    try:
        query = get_student_placement_query(db, dept_id, batch_id, placement_status, company_id)
        results = query.all()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=16,
            textColor=colors.HexColor("#1F4E79"),
            alignment=1, # Center
            spaceAfter=20
        )
        
        story = [
            Paragraph("Student Placement Status Report", title_style),
            Spacer(1, 10)
        ]

        # Table headers and rows
        data = [["USN", "Name", "Branch", "CGPA", "Company", "Designation", "CTC", "Status"]]
        for r in results:
            data.append([
                r.usno,
                r.name,
                r.branch,
                str(r.cgpa) if r.cgpa else "0.0",
                r.company or "-",
                r.designation or "-",
                f"{float(r.ctc)} LPA" if r.ctc else "-",
                r.status
            ])

        # Render table with reportlab
        t = Table(data, colWidths=[90, 110, 60, 50, 100, 100, 70, 70])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1F4E79")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('BACKGROUND', (0,1), (-1,-1), colors.HexColor("#F2F4F7")),
            ('GRID', (0,0), (-1,-1), 1, colors.HexColor("#D1D5DB")),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 9),
            ('BOTTOMPADDING', (0,1), (-1,-1), 6),
            ('TOPPADDING', (0,1), (-1,-1), 6),
        ]))
        story.append(t)
        doc.build(story)

        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=Student_Placement_Report.pdf"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# -------------------------------------------------------------
# 5. Get Drive Summary Report Endpoint
# -------------------------------------------------------------
@router.get("/drive-summary")
def get_drive_summary_report(
    company_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    try:
        # Build filter conditions
        query = db.query(PlacementDrive)
        if company_id:
            query = query.filter(PlacementDrive.company_id == company_id)
        if status and status != "All":
            # Map frontend labels (Ongoing, Completed, Upcoming) to DB status enums or values
            # DB values are: 0=Draft, 1=Scheduled, 2=Active, 3=Closed, 4=Cancelled
            # Map: Completed -> [3], Ongoing -> [2], Upcoming -> [1]
            if status == "Completed":
                query = query.filter(PlacementDrive.status == 3)
            elif status == "Ongoing":
                query = query.filter(PlacementDrive.status == 2)
            elif status == "Upcoming":
                query = query.filter(PlacementDrive.status == 1)
        if date_from:
            query = query.filter(PlacementDrive.drive_date >= date_from)
        if date_to:
            query = query.filter(PlacementDrive.drive_date <= date_to)

        drives = query.all()
        report_data = []

        for d in drives:
            # Query company name directly
            company_name = db.query(PlacementCompany.company_name).filter(PlacementCompany.company_id == d.company_id).scalar() or "-"

            # Gather dynamic real-time funnel counts from database
            applications_count = db.query(func.count(PlacementApplication.application_id))\
                .filter(PlacementApplication.drive_id == d.drive_id).scalar() or 0
            
            eligible_count = d.eligible_student_count or db.query(func.count(PLMStudentProfile.profile_id))\
                .filter(PLMStudentProfile.is_placement_eligible == 1).scalar() or 0
            
            offers_issued = db.query(func.count(PlacementOffer.offer_id))\
                .filter(PlacementOffer.drive_id == d.drive_id).scalar() or 0
                
            offers_accepted = db.query(func.count(PlacementOffer.offer_id))\
                .filter(PlacementOffer.drive_id == d.drive_id, PlacementOffer.status == 'ACCEPTED').scalar() or 0

            # Real shortlist count from plm_shortlist table or derived from offers issued
            shortlisted_from_db = db.query(func.count(PlacementShortlist.shortlist_id))\
                .join(PlacementApplication, PlacementShortlist.application_id == PlacementApplication.application_id)\
                .filter(PlacementApplication.drive_id == d.drive_id).scalar() or 0
            
            shortlisted_count = max(shortlisted_from_db, offers_issued, offers_accepted)

            waitlisted_count = db.query(func.count(PlacementWaitlist.waitlist_id))\
                .filter(PlacementWaitlist.drive_id == d.drive_id).scalar() or 0
            
            interviewed_from_db = db.query(func.count(PlacementApplication.application_id))\
                .filter(PlacementApplication.drive_id == d.drive_id, PlacementApplication.status == 'INTERVIEW').scalar() or 0
            
            interviewed_count = max(interviewed_from_db, shortlisted_count)

            rejected_from_db = db.query(func.count(PlacementApplication.application_id))\
                .filter(PlacementApplication.drive_id == d.drive_id, PlacementApplication.status == 'REJECTED').scalar() or 0
            
            rejected_count = max(0, applications_count - shortlisted_count) if rejected_from_db == 0 else rejected_from_db

            report_data.append({
                "drive_id": d.drive_id,
                "drive_name": d.drive_name,
                "company_name": company_name,
                "applications": applications_count,
                "eligible": eligible_count,
                "shortlisted": shortlisted_count,
                "waitlisted": waitlisted_count,
                "rejected": rejected_count,
                "interviewed": interviewed_count,
                "offers_issued": offers_issued,
                "offers_accepted": offers_accepted
            })

        return returnSuccess(report_data, "Drive summary report fetched successfully")
    except Exception as e:
        return returnException(str(e))


# -------------------------------------------------------------
# 6. Export Drive Summary Excel Endpoint
# -------------------------------------------------------------
@router.get("/drive-summary/export-excel")
def export_drive_summary_excel(
    company_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    try:
        # Build filter conditions
        query = db.query(PlacementDrive)
        if company_id:
            query = query.filter(PlacementDrive.company_id == company_id)
        if status and status != "All":
            if status == "Completed":
                query = query.filter(PlacementDrive.status == 3)
            elif status == "Ongoing":
                query = query.filter(PlacementDrive.status == 2)
            elif status == "Upcoming":
                query = query.filter(PlacementDrive.status == 1)
        if date_from:
            query = query.filter(PlacementDrive.drive_date >= date_from)
        if date_to:
            query = query.filter(PlacementDrive.drive_date <= date_to)

        drives = query.all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Drive Summary"

        # Headers
        headers = [
            "Drive Name", "Company", "Applications", "Eligible", 
            "Shortlisted", "Waitlisted", "Rejected", "Interviewed", 
            "Offers Issued", "Offers Accepted"
        ]
        ws.append(headers)

        # Style headers
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Rows
        for d in drives:
            company_name = db.query(PlacementCompany.company_name).filter(PlacementCompany.company_id == d.company_id).scalar() or "-"
            
            applications_count = db.query(func.count(PlacementApplication.application_id))\
                .filter(PlacementApplication.drive_id == d.drive_id).scalar() or 0
            
            eligible_count = d.eligible_student_count or db.query(func.count(PLMStudentProfile.profile_id))\
                .filter(PLMStudentProfile.is_placement_eligible == 1).scalar() or 0
            
            offers_issued = db.query(func.count(PlacementOffer.offer_id))\
                .filter(PlacementOffer.drive_id == d.drive_id).scalar() or 0
                
            offers_accepted = db.query(func.count(PlacementOffer.offer_id))\
                .filter(PlacementOffer.drive_id == d.drive_id, PlacementOffer.status == 'ACCEPTED').scalar() or 0

            shortlisted_from_db = db.query(func.count(PlacementShortlist.shortlist_id))\
                .join(PlacementApplication, PlacementShortlist.application_id == PlacementApplication.application_id)\
                .filter(PlacementApplication.drive_id == d.drive_id).scalar() or 0
            
            shortlisted_count = max(shortlisted_from_db, offers_issued, offers_accepted)

            waitlisted_count = db.query(func.count(PlacementWaitlist.waitlist_id))\
                .filter(PlacementWaitlist.drive_id == d.drive_id).scalar() or 0
            
            interviewed_from_db = db.query(func.count(PlacementApplication.application_id))\
                .filter(PlacementApplication.drive_id == d.drive_id, PlacementApplication.status == 'INTERVIEW').scalar() or 0
            
            interviewed_count = max(interviewed_from_db, shortlisted_count)

            rejected_from_db = db.query(func.count(PlacementApplication.application_id))\
                .filter(PlacementApplication.drive_id == d.drive_id, PlacementApplication.status == 'REJECTED').scalar() or 0
            
            rejected_count = max(0, applications_count - shortlisted_count) if rejected_from_db == 0 else rejected_from_db

            ws.append([
                d.drive_name,
                company_name,
                applications_count,
                eligible_count,
                shortlisted_count,
                waitlisted_count,
                rejected_count,
                interviewed_count,
                offers_issued,
                offers_accepted
            ])

        # Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)

        return StreamingResponse(
            file_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=Recruiter_Drive_Summary_Report.xlsx"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# -------------------------------------------------------------
# 7. Export Drive Summary PDF Endpoint
# -------------------------------------------------------------
@router.get("/drive-summary/export-pdf")
def export_drive_summary_pdf(
    company_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    try:
        # Build filter conditions
        query = db.query(PlacementDrive)
        if company_id:
            query = query.filter(PlacementDrive.company_id == company_id)
        if status and status != "All":
            if status == "Completed":
                query = query.filter(PlacementDrive.status == 3)
            elif status == "Ongoing":
                query = query.filter(PlacementDrive.status == 2)
            elif status == "Upcoming":
                query = query.filter(PlacementDrive.status == 1)
        if date_from:
            query = query.filter(PlacementDrive.drive_date >= date_from)
        if date_to:
            query = query.filter(PlacementDrive.drive_date <= date_to)

        drives = query.all()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=16,
            textColor=colors.HexColor("#1F4E79"),
            alignment=1, # Center
            spaceAfter=20
        )
        
        story = [
            Paragraph("Recruiter Drive Summary Report", title_style),
            Spacer(1, 10)
        ]

        # Table headers and rows
        data = [[
            "Drive Name", "Company", "App.", "Elig.", 
            "Short.", "Wait.", "Rej.", "Interv.", 
            "Issued", "Acc."
        ]]
        
        for d in drives:
            company_name = db.query(PlacementCompany.company_name).filter(PlacementCompany.company_id == d.company_id).scalar() or "-"
            
            applications_count = db.query(func.count(PlacementApplication.application_id))\
                .filter(PlacementApplication.drive_id == d.drive_id).scalar() or 0
            
            eligible_count = d.eligible_student_count or db.query(func.count(PLMStudentProfile.profile_id))\
                .filter(PLMStudentProfile.is_placement_eligible == 1).scalar() or 0
            
            offers_issued = db.query(func.count(PlacementOffer.offer_id))\
                .filter(PlacementOffer.drive_id == d.drive_id).scalar() or 0
                
            offers_accepted = db.query(func.count(PlacementOffer.offer_id))\
                .filter(PlacementOffer.drive_id == d.drive_id, PlacementOffer.status == 'ACCEPTED').scalar() or 0

            shortlisted_from_db = db.query(func.count(PlacementShortlist.shortlist_id))\
                .join(PlacementApplication, PlacementShortlist.application_id == PlacementApplication.application_id)\
                .filter(PlacementApplication.drive_id == d.drive_id).scalar() or 0
            
            shortlisted_count = max(shortlisted_from_db, offers_issued, offers_accepted)

            waitlisted_count = db.query(func.count(PlacementWaitlist.waitlist_id))\
                .filter(PlacementWaitlist.drive_id == d.drive_id).scalar() or 0
            
            interviewed_from_db = db.query(func.count(PlacementApplication.application_id))\
                .filter(PlacementApplication.drive_id == d.drive_id, PlacementApplication.status == 'INTERVIEW').scalar() or 0
            
            interviewed_count = max(interviewed_from_db, shortlisted_count)

            rejected_from_db = db.query(func.count(PlacementApplication.application_id))\
                .filter(PlacementApplication.drive_id == d.drive_id, PlacementApplication.status == 'REJECTED').scalar() or 0
            
            rejected_count = max(0, applications_count - shortlisted_count) if rejected_from_db == 0 else rejected_from_db

            data.append([
                d.drive_name,
                company_name,
                str(applications_count),
                str(eligible_count),
                str(shortlisted_count),
                str(waitlisted_count),
                str(rejected_count),
                str(interviewed_count),
                str(offers_issued),
                str(offers_accepted)
            ])

        # Widths (total 700 width for landscape page)
        t = Table(data, colWidths=[130, 110, 50, 50, 60, 60, 60, 60, 60, 60])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1F4E79")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 9),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('BACKGROUND', (0,1), (-1,-1), colors.HexColor("#F2F4F7")),
            ('GRID', (0,0), (-1,-1), 1, colors.HexColor("#D1D5DB")),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 9),
            ('BOTTOMPADDING', (0,1), (-1,-1), 6),
            ('TOPPADDING', (0,1), (-1,-1), 6),
        ]))
        story.append(t)
        doc.build(story)

        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=Recruiter_Drive_Summary_Report.pdf"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
